from pathlib import Path

from docx import Document

from app.ingestion.models.document import DocumentPage, ExtractedDocument


def extract_docx(path: Path) -> ExtractedDocument:
    """
    Extract text from a DOCX document.
    """
    document = Document(str(path))

    paragraphs = [
        paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()
    ]

    text = "\n".join(paragraphs)

    return ExtractedDocument(
        filename=path.name,
        content_type=("application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        pages=[
            DocumentPage(
                page_number=1,
                text=text,
            )
        ],
    )
